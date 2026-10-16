
import os
import re
import sqlite3
import secrets
import smtplib
from datetime import datetime
from io import BytesIO
from email.utils import parseaddr

from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    session,
    jsonify,
    send_file,
    abort
)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


# --------------------------------------------------
# APP CONFIGURATION
# --------------------------------------------------

app = Flask(__name__)
app.secret_key = os.environ.get(
    "FLASK_SECRET_KEY", "change-this-secret-key"
)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATABASE = os.path.join(BASE_DIR, "queue.db")

ADMIN_USERNAME = os.environ.get("ADMIN_USERNAME", "admin")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "admin123")


# --------------------------------------------------
# DATABASE CONNECTION
# --------------------------------------------------

def get_db():
    conn = sqlite3.connect(DATABASE)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA busy_timeout = 5000")
    return conn


def init_db():
    conn = get_db()

    conn.execute("""
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            token TEXT UNIQUE,
            name TEXT NOT NULL,
            mobile TEXT NOT NULL,
            email TEXT,
            service TEXT NOT NULL DEFAULT 'General Enquiry',
            people INTEGER DEFAULT 1,
            priority TEXT DEFAULT 'Normal',
            preferred_time TEXT,
            message TEXT,
            status TEXT DEFAULT 'Waiting',
            counter TEXT DEFAULT 'Counter 1',
            created_at TEXT,
            place TEXT DEFAULT '',
            category TEXT DEFAULT '',
            reason TEXT DEFAULT ''
        )
    """)

    # Add new columns to an existing database safely.
    existing_columns = {
        row["name"]
        for row in conn.execute(
            "PRAGMA table_info(users)"
        ).fetchall()
    }

    new_columns = {
        "place": "TEXT DEFAULT ''",
        "category": "TEXT DEFAULT ''",
        "reason": "TEXT DEFAULT ''"
    }

    for column, data_type in new_columns.items():
        if column not in existing_columns:
            conn.execute(
                f"ALTER TABLE users ADD COLUMN {column} {data_type}"
            )

    # Update older rows with a missing timestamp.
    conn.execute("""
        UPDATE users
        SET created_at = ?
        WHERE created_at IS NULL OR created_at = ''
    """, (datetime.now().strftime("%Y-%m-%d %H:%M:%S"),))

    conn.commit()
    conn.close()


# --------------------------------------------------
# HELPER FUNCTIONS
# --------------------------------------------------

def admin_required():
    return bool(session.get("admin"))


def priority_order():
    return """
        CASE priority
            WHEN 'Emergency' THEN 1
            WHEN 'Senior Citizen' THEN 2
            ELSE 3
        END,
        id ASC
    """


def make_token():
    return "Q" + secrets.token_hex(3).upper()


def valid_email(email):
    if not email:
        return True

    if len(email) > 254 or any(c.isspace() for c in email):
        return False

    _, parsed_email = parseaddr(email)

    return (
        parsed_email == email
        and re.fullmatch(
            r"[^@\s]+@[^@\s]+\.[^@\s]+", email
        ) is not None
    )


# --------------------------------------------------
# HOME PAGE
# --------------------------------------------------

@app.route("/")
def home():
    return render_template("index.html")


# --------------------------------------------------
# SELECT PLACE PAGE
# --------------------------------------------------

@app.route("/select-place")
def select_place():
    return render_template("select-place.html")


# --------------------------------------------------
# JOIN QUEUE
# --------------------------------------------------

@app.route("/join-queue", methods=["GET", "POST"])
def join_queue():

    if request.method == "GET":
        return render_template("join-queue.html")

    # Accept common field-name variations from HTML forms.
    name = request.form.get("name", "").strip()
    mobile = request.form.get("mobile", "").strip()
    email = request.form.get("email", "").strip()

    place = (
        request.form.get("place")
        or request.form.get("selected_place")
        or request.form.get("hospital")
        or ""
    ).strip()

    category = (
        request.form.get("category")
        or request.form.get("selected_category")
        or request.form.get("place_category")
        or ""
    ).strip()

    reason = (
        request.form.get("reason")
        or request.form.get("visit_reason")
        or request.form.get("service")
        or category
        or "General Enquiry"
    ).strip()

    people_value = request.form.get("people", "1").strip()
    priority = request.form.get("priority", "Normal").strip()
    preferred_time = request.form.get("preferred_time", "").strip()
    message = request.form.get("message", "").strip()

    # Required-field validation.
    if not name or not mobile:
        return "Name and mobile number are required.", 400

    # Exactly 10 numeric digits.
    if not re.fullmatch(r"[0-9]{10}", mobile):
        return "Enter a valid 10-digit mobile number.", 400

    if not valid_email(email):
        return "Enter a valid email address.", 400

    try:
        people = int(people_value)
        if people < 1 or people > 50:
            raise ValueError
    except ValueError:
        return "Enter a valid number of people.", 400

    allowed_priorities = [
        "Normal",
        "Senior Citizen",
        "Emergency"
    ]

    if priority not in allowed_priorities:
        priority = "Normal"

    # The selected category becomes the service.
    # This prevents a hospital visit from always becoming
    # General Enquiry.
    service = category or reason or "General Enquiry"

    token = make_token()
    created_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    conn = get_db()

    try:
        # Retry if the generated token already exists.
        for attempt in range(5):
            try:
                conn.execute("""
                    INSERT INTO users (
                        token, name, mobile, email, service,
                        people, priority, preferred_time,
                        message, status, counter, created_at,
                        place, category, reason
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    token,
                    name,
                    mobile,
                    email,
                    service,
                    people,
                    priority,
                    preferred_time,
                    message,
                    "Waiting",
                    "Counter 1",
                    created_at,
                    place,
                    category,
                    reason
                ))

                conn.commit()
                break

            except sqlite3.IntegrityError:
                if attempt == 4:
                    raise

                token = make_token()

        return redirect(url_for("token_page", token=token))

    except sqlite3.Error:
        conn.rollback()
        app.logger.exception("Unable to save queue entry")
        return "Unable to join the queue. Please try again.", 500

    finally:
        conn.close()


# --------------------------------------------------
# TOKEN PAGE
# --------------------------------------------------

@app.route("/token/<token>")
def token_page(token):
    conn = get_db()

    user = conn.execute(
        "SELECT * FROM users WHERE token = ?",
        (token,)
    ).fetchone()

    conn.close()

    if user is None:
        abort(404)

    return render_template("token.html", user=user)


# --------------------------------------------------
# QUEUE STATUS
# --------------------------------------------------

@app.route("/queue-status")
def queue_status():
    token = request.args.get("token", "").strip()

    if not token:
        return render_template("queue-status.html", user=None)

    conn = get_db()

    user = conn.execute(
        "SELECT * FROM users WHERE token = ?",
        (token,)
    ).fetchone()

    conn.close()

    return render_template("queue-status.html", user=user)


# --------------------------------------------------
# ADMIN LOGIN
# --------------------------------------------------

@app.route("/admin-login", methods=["GET", "POST"])
@app.route("/admin.html", methods=["GET", "POST"])
def admin_login():

    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")

        if (
            username == ADMIN_USERNAME
            and password == ADMIN_PASSWORD
        ):
            session.clear()
            session["admin"] = True
            return redirect(url_for("admin"))

        return render_template(
            "admin-login.html",
            error="Invalid username or password."
        )

    if admin_required():
        return redirect(url_for("admin"))

    return render_template("admin-login.html")


# --------------------------------------------------
# ADMIN DASHBOARD
# --------------------------------------------------

@app.route("/admin")
@app.route("/admin.html/dashboard")
def admin():

    if not admin_required():
        return redirect(url_for("admin_login"))

    conn = get_db()

    users = conn.execute("""
        SELECT * FROM users
        ORDER BY
            CASE status
                WHEN 'Serving' THEN 1
                WHEN 'Waiting' THEN 2
                WHEN 'Skipped' THEN 3
                WHEN 'Completed' THEN 4
                ELSE 5
            END,
            CASE priority
                WHEN 'Emergency' THEN 1
                WHEN 'Senior Citizen' THEN 2
                ELSE 3
            END,
            id ASC
    """).fetchall()

    serving = conn.execute("""
        SELECT * FROM users
        WHERE status = 'Serving'
        ORDER BY id DESC
        LIMIT 1
    """).fetchone()

    waiting_count = conn.execute("""
        SELECT COUNT(*) FROM users
        WHERE status = 'Waiting'
    """).fetchone()[0]

    completed_count = conn.execute("""
        SELECT COUNT(*) FROM users
        WHERE status = 'Completed'
    """).fetchone()[0]

    conn.close()

    return render_template(
        "admin.html",
        users=users,
        queue=users,
        serving=serving,
        waiting_count=waiting_count,
        completed_count=completed_count
    )


# --------------------------------------------------
# CALL NEXT PERSON
# --------------------------------------------------

@app.route("/call-next", methods=["POST", "GET"])
@app.route("/call-next/<int:user_id>", methods=["POST", "GET"])
def call_next(user_id=None):

    if not admin_required():
        return redirect(url_for("admin_login"))

    conn = get_db()

    try:
        if user_id is not None:
            user = conn.execute("""
                SELECT * FROM users
                WHERE id = ? AND status = 'Waiting'
            """, (user_id,)).fetchone()
        else:
            user = conn.execute(f"""
                SELECT * FROM users
                WHERE status = 'Waiting'
                ORDER BY {priority_order()}
                LIMIT 1
            """).fetchone()

        if user:
            conn.execute("""
                UPDATE users
                SET status = 'Serving'
                WHERE id = ?
            """, (user["id"],))

            conn.commit()

    finally:
        conn.close()

    return redirect(url_for("admin"))


# --------------------------------------------------
# COMPLETE TOKEN
# --------------------------------------------------

@app.route("/complete/<int:user_id>", methods=["POST", "GET"])
def complete(user_id):

    if not admin_required():
        return redirect(url_for("admin_login"))

    conn = get_db()

    conn.execute("""
        UPDATE users
        SET status = 'Completed'
        WHERE id = ? AND status = 'Serving'
    """, (user_id,))

    conn.commit()
    conn.close()

    return redirect(url_for("admin"))


# --------------------------------------------------
# SKIP TOKEN
# --------------------------------------------------

@app.route("/skip/<int:user_id>", methods=["POST", "GET"])
def skip(user_id):

    if not admin_required():
        return redirect(url_for("admin_login"))

    conn = get_db()

    conn.execute("""
        UPDATE users
        SET status = 'Skipped'
        WHERE id = ? AND status IN ('Waiting', 'Serving')
    """, (user_id,))

    conn.commit()
    conn.close()

    return redirect(url_for("admin"))


# --------------------------------------------------
# QUEUE INFORMATION / COUNTERS
# --------------------------------------------------

@app.route("/queue-info")
def queue_info():

    conn = get_db()

    total = conn.execute(
        "SELECT COUNT(*) FROM users"
    ).fetchone()[0]

    waiting = conn.execute("""
        SELECT COUNT(*) FROM users
        WHERE status = 'Waiting'
    """).fetchone()[0]

    completed = conn.execute("""
        SELECT COUNT(*) FROM users
        WHERE status = 'Completed'
    """).fetchone()[0]

    emergency = conn.execute("""
        SELECT COUNT(*) FROM users
        WHERE status = 'Waiting' AND priority = 'Emergency'
    """).fetchone()[0]

    senior_citizen = conn.execute("""
        SELECT COUNT(*) FROM users
        WHERE status = 'Waiting'
          AND priority = 'Senior Citizen'
    """).fetchone()[0]

    # Counts by selected service/category.
    service_counts = conn.execute("""
        SELECT service, COUNT(*) AS count
        FROM users
        WHERE status = 'Waiting'
        GROUP BY service
        ORDER BY service
    """).fetchall()

    # Counts by counter.
    counter_counts = conn.execute("""
        SELECT counter, COUNT(*) AS count
        FROM users
        WHERE status = 'Waiting'
        GROUP BY counter
        ORDER BY counter
    """).fetchall()

    average_wait = waiting * 5

    conn.close()

    return jsonify({
        "total": total,
        "waiting": waiting,
        "completed": completed,
        "emergency": emergency,
        "senior_citizen": senior_citizen,
        "average_wait": average_wait,
        "service_counts": {
            row["service"]: row["count"]
            for row in service_counts
        },
        "counter_counts": {
            row["counter"]: row["count"]
            for row in counter_counts
        }
    })


# --------------------------------------------------
# ABOUT PAGE
# --------------------------------------------------

@app.route("/about")
def about():
    return render_template("about.html")


# --------------------------------------------------
# LOGOUT
# --------------------------------------------------

@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("home"))


# --------------------------------------------------
# EXPORT DATABASE TO EXCEL
# --------------------------------------------------

@app.route("/export-excel")
def export_excel():

    # Only the admin can download queue records.
    if not admin_required():
        return redirect(url_for("admin_login"))

    conn = get_db()

    try:
        cursor = conn.execute("""
            SELECT * FROM users
            ORDER BY id ASC
        """)

        rows = cursor.fetchall()
        columns = [description[0] for description in cursor.description]

    finally:
        conn.close()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Queue Records"

    # Write headings.
    sheet.append(columns)

    # Write database records.
    for row in rows:
        sheet.append([row[column] for column in columns])

    # Style heading row.
    for cell in sheet[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1565C0"
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    # Adjust column widths.
    for column in sheet.columns:
        max_length = max(
            len(str(cell.value or ""))
            for cell in column
        )

        sheet.column_dimensions[
            column[0].column_letter
        ].width = min(max_length + 3, 35)

    # Create the Excel file in memory.
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    filename = (
        "queue_data_"
        + datetime.now().strftime("%Y%m%d_%H%M%S")
        + ".xlsx"
    )

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet"
        )
    )


# --------------------------------------------------
# START APPLICATION
# --------------------------------------------------

init_db()

if __name__ == "__main__":
    app.run(
        host="127.0.0.1",
        port=2007,
        debug=True
    )